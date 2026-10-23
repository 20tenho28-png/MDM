#!/usr/bin/env python3
"""Reconstrói o percurso de cada pedido da MDM a partir dos canais em bruto.

Entrada (pasta --entrada, ver 00-exportar-dados.md):
  *.mbox                     emails recebidos (pedidos): gmail.mbox do Takeout e imap-comercial.mbox de exportar-imap.py
  remetentes-excluir.txt     emails ou domínios que não são clientes (fornecedores, bancos, a própria MDM), um por linha
  whatsapp/**/*.txt          conversas exportadas sem multimédia (as do iPhone chamam-se _chat.txt: uma pasta por conversa,
                             com o nome do cliente)
  chamadas-*.csv             data,hora,numero,duracao_s,tipo
  agenda.csv                 visitas marcadas (colunas com data e cliente, detetadas por nome)
  orc/*.xlsx                 orçamentos (cliente do nome do ficheiro; data e valor da primeira folha)
  saft/*.xml                 SAF-T mensal (faturas, notas de crédito, recibos); os meses depois do período só dão pagamentos
  pedidos.csv                export do registo (exportar-supabase.sql)

Saída (pasta --saida):
  pedidos_reconstruidos.csv  uma linha por pedido (um cliente pode ter vários), uma coluna por etapa (data ou vazio)
  funil.md                   contagens, conversão por etapa, medianas de dias, euros
  reconciliacao.md           faturas, notas de crédito e líquido do SAF-T por mês vs. o que o percurso encontrou
  chamadas.md                recebidas e perdidas por mês e por linha
  nao_correspondidos.csv     o que não casou com nada (para corrigir à mão)
  remetentes.csv             quantos emails por domínio, para montar remetentes-excluir.txt
  orc_sem_data.csv           orçamentos sem data na folha (ficam fora do funil até se corrigir)

Uso:
  python3 reconstruir.py --entrada entrada --saida saida --desde 2026-07-01 --ate 2026-09-30
  python3 reconstruir.py --demo        # gera dados fictícios e corre, para ver o formato

Só biblioteca padrão mais openpyxl (pip install openpyxl). Nunca envia nada para a rede.
"""
import argparse, csv, glob, mailbox, os, re, statistics, sys, unicodedata, datetime as dt
from collections import Counter
from email.header import decode_header, make_header
from email.utils import parsedate_to_datetime, parseaddr
from xml.etree import ElementTree as ET

ETAPAS = ["pedido", "visita", "orcamento", "adjudicacao", "obra", "fatura", "pagamento"]
STOP = {"lda", "ltd", "sa", "unipessoal", "e", "de", "da", "do", "dos", "das", "sr", "sra", "eng", "dr", "dra", "n", "no", "nº"}
# Palavras que muitos clientes partilham: sozinhas não chegam para dizer que dois nomes são o mesmo cliente.
GENERICOS = {"condominio", "escola", "basica", "secundaria", "clinica", "dentaria", "restaurante", "hotel", "farmacia", "ginasio",
             "rua", "avenida", "av", "praca", "largo", "lisboa", "loures", "amadora", "oeiras", "lote", "predio", "edificio",
             "construcoes", "industria", "banco", "particular", "empresa", "sociedade", "servicos", "comercio", "obras", "centro"}
DOMINIOS_MDM = {"mdmassist.pt", "mdmassist.com"}
EMAILS_MDM = {"mdmassistencia@gmail.com", "mdmassist@mdmassist.com"}

# ---------- utilidades ----------
def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    # os números ficam sempre (n.º 3 e n.º 5 são escolas diferentes); as palavras de uma letra não
    return " ".join(t for t in s.split() if t not in STOP and (t.isdigit() or len(t) > 1))

def parece_mesmo(a, b):
    ta, tb = set(norm(a).split()), set(norm(b).split())
    if not ta or not tb: return False
    na, nb = {t for t in ta if t.isdigit()}, {t for t in tb if t.isdigit()}
    if na and nb and na != nb: return False                     # números diferentes: clientes diferentes
    comuns = ta & tb
    fortes = comuns - GENERICOS
    if len(ta) == 1 or len(tb) == 1:                            # nome de uma palavra só: tem de ser essa palavra e não genérica
        return bool(fortes)
    return len(comuns) >= 2 and len(fortes) >= 1

def data_iso(x):
    if x is None: return None
    if isinstance(x, (dt.datetime, dt.date)): return x.strftime("%Y-%m-%d")
    s = str(x).strip()
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m: return "-".join(m.groups())
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", s) or re.search(r"(\d{1,2})-(\d{1,2})-(\d{4})", s)
    if m: return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{2})\b", s)
    if m: return f"20{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return None

def no_periodo(d, desde, ate):
    return d is not None and desde <= d <= ate

def dias(a, b):
    if not a or not b: return None
    return (dt.date.fromisoformat(b) - dt.date.fromisoformat(a)).days

def mediana(xs):
    xs = [x for x in xs if x is not None]
    return round(statistics.median(xs), 1) if xs else None

def eur(v):
    """1 234,56 € (como na MDM), nunca 1,234.56."""
    s = f"{abs(v):,.2f}".replace(",", " ").replace(".", ",")
    return ("-" if v < 0 else "") + s + " €"

def num1(v):
    return f"{v}".replace(".", ",") if v is not None else "sem dados"

def _num(x):
    """Lê 1180.5 (export do registo), 1 180,50 e 1.180,50 (escrito à mão) como o mesmo número."""
    if x in (None, ""): return None
    s = str(x).strip().replace(" ", "").replace("€", "")
    if "," in s: s = s.replace(".", "").replace(",", ".")
    try: return float(s)
    except ValueError: return None

def ler_csv(caminho):
    if not os.path.exists(caminho): return []
    with open(caminho, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))

def coluna(row, *cands):
    for k in row:
        kn = norm(k)
        if any(c in kn for c in cands): return row[k]
    return None

def gravar(caminho, texto):
    with open(caminho, "w", encoding="utf-8") as f:
        f.write(texto)

# ---------- leitores por fonte ----------
def ler_excluir(entrada):
    caminho = os.path.join(entrada, "remetentes-excluir.txt")
    if not os.path.exists(caminho): return set()
    with open(caminho, encoding="utf-8") as f:
        return {l.strip().lower() for l in f if l.strip() and not l.startswith("#")}

def excluido(email, excluir):
    email = (email or "").lower()
    dom = email.split("@")[-1] if "@" in email else ""
    return email in EMAILS_MDM or dom in DOMINIOS_MDM or email in excluir or dom in excluir or ("@" + dom) in excluir

def ler_emails(entrada, desde, ate, toques, meta, excluir):
    n, fora, dominios = 0, 0, Counter()
    for mb in sorted(glob.glob(os.path.join(entrada, "*.mbox"))):
        for msg in mailbox.mbox(mb):
            try: d = parsedate_to_datetime(msg["Date"]).strftime("%Y-%m-%d")
            except Exception: continue
            if not no_periodo(d, desde, ate): continue
            de = str(make_header(decode_header(msg.get("From", ""))))
            assunto = str(make_header(decode_header(msg.get("Subject", ""))))
            nome, email = parseaddr(de)
            dominios[email.split("@")[-1].lower() if "@" in email else "(sem endereço)"] += 1
            if excluido(email, excluir) or re.search(r"newsletter|unsubscribe|promo|noreply|no-reply|mailer-daemon", (de + assunto).lower()):
                fora += 1; continue
            toques.append({"fonte": "email", "data": d, "cliente": nome or email, "detalhe": assunto[:80], "etapa": "pedido"})
            n += 1
    meta["emails"], meta["emails_excluidos"], meta["dominios"] = n, fora, dominios

def nome_whatsapp(f, linhas_autores, eu):
    base = os.path.splitext(os.path.basename(f))[0]
    for pre in ("WhatsApp Chat with", "WhatsApp Chat -", "Conversa do WhatsApp com", "Conversa do WhatsApp -"):
        base = base.replace(pre, "")
    base = base.strip()
    if base and not base.startswith("_chat"): return base
    pasta = os.path.basename(os.path.dirname(f))
    for pre in ("WhatsApp Chat -", "Conversa do WhatsApp -"):
        pasta = pasta.replace(pre, "")
    if pasta.strip() and pasta != "whatsapp": return pasta.strip()
    # sem nome no ficheiro nem na pasta: quem mais escreve e não é da MDM
    outros = Counter(a for a in linhas_autores if not any(norm(e) and norm(e) in norm(a) for e in eu))
    return outros.most_common(1)[0][0] if outros else base

def ler_whatsapp(entrada, desde, ate, toques, eu):
    pad = re.compile(r"^\[?(\d{1,2}/\d{1,2}/\d{2,4})[, ]+\d{1,2}:\d{2}(?::\d{2})?\]?\s*[-–]?\s*([^:]+):")
    for f in sorted(glob.glob(os.path.join(entrada, "whatsapp", "**", "*.txt"), recursive=True)):
        primeira, autores = None, []
        with open(f, encoding="utf-8", errors="ignore") as fh:
            for linha in fh:
                m = pad.match(linha.strip("‎﻿"))
                if not m: continue
                autores.append(m.group(2).strip())
                d = data_iso(m.group(1))
                if no_periodo(d, desde, ate) and primeira is None: primeira = d
        if primeira:
            toques.append({"fonte": "whatsapp", "data": primeira, "cliente": nome_whatsapp(f, autores, eu), "detalhe": "primeira mensagem no período", "etapa": "pedido"})

def ler_chamadas(entrada, desde, ate):
    rows = []
    for f in glob.glob(os.path.join(entrada, "chamadas-*.csv")):
        for r in ler_csv(f):
            d = data_iso(coluna(r, "data"))
            if not no_periodo(d, desde, ate): continue
            tipo = norm(coluna(r, "tipo") or "")
            dur = coluna(r, "dura") or "0"
            try: dur = float(str(dur).replace(",", "."))
            except ValueError: dur = 0
            perdida = "perd" in tipo or "missed" in tipo or ("receb" in tipo and dur == 0)
            rows.append({"data": d, "mes": d[:7], "linha": os.path.basename(f), "recebida": "receb" in tipo or perdida, "perdida": perdida})
    return rows

def ler_agenda(entrada, desde, ate, toques):
    for nome in ("agenda.csv", "agenda-outlook.csv"):
        for r in ler_csv(os.path.join(entrada, nome)):
            d = data_iso(coluna(r, "data", "start", "inicio"))
            cli = coluna(r, "cliente", "nome", "subject", "assunto", "local")
            if no_periodo(d, desde, ate) and cli:
                toques.append({"fonte": nome, "data": d, "cliente": cli, "detalhe": "visita", "etapa": "visita"})
    for r in ler_csv(os.path.join(entrada, "folhas-obra.csv")):
        d = data_iso(coluna(r, "data"))
        cli = coluna(r, "cliente", "nome")
        if no_periodo(d, desde, ate) and cli:
            toques.append({"fonte": "folhas-obra", "data": d, "cliente": cli, "detalhe": coluna(r, "ob", "fo", "num") or "obra", "etapa": "obra"})

def data_do_orc(linhas):
    """A data ao lado de «Data»; senão a primeira data da folha que não seja a de validade. Nunca a do ficheiro no disco."""
    primeira = None
    for row in linhas:
        for i, c in enumerate(row):
            if isinstance(c, str) and re.match(r"^\s*data\b", c.lower()) and "valid" not in c.lower():
                for v in row[i + 1:i + 4]:
                    if isinstance(v, (dt.datetime, dt.date)): return data_iso(v)
                d = data_iso(c)
                if d: return d
            if isinstance(c, (dt.datetime, dt.date)) and primeira is None:
                antes = " ".join(str(x).lower() for x in row[:i] if isinstance(x, str))
                if "valid" not in antes: primeira = data_iso(c)
    return primeira

def ler_orc(entrada, desde, ate, toques, sem_data):
    try: import openpyxl
    except ImportError:
        print("aviso: openpyxl em falta (pip install openpyxl); ORC ignorados", file=sys.stderr); return
    for f in sorted(glob.glob(os.path.join(entrada, "orc", "*.xlsx"))):
        base = os.path.splitext(os.path.basename(f))[0]
        m = re.match(r"(ORC\d+_\d{4})\s*[-–]?\s*(.*)", base)
        num, cli = (m.group(1), m.group(2)) if m else (base, base)
        d, valor = None, None
        try:
            wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
            linhas = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(min_row=1, max_row=60, values_only=True)]
            wb.close()
            d = data_do_orc(linhas)
            for row in linhas:
                for c in row:
                    if isinstance(c, str) and "valor total" in c.lower():
                        mv = re.search(r"([\d\.\s]+,\d{2})\s*€", c)
                        if mv: valor = _num(mv.group(1))
        except Exception as e:
            print(f"aviso: não li {f}: {e}", file=sys.stderr)
        if d is None:
            d = data_iso(base)                     # uma data escrita no nome do ficheiro
        if d is None:
            sem_data.append({"ficheiro": os.path.basename(f), "num": num}); continue
        if no_periodo(d, desde, ate):
            toques.append({"fonte": "orc", "data": d, "cliente": cli, "detalhe": num, "etapa": "orcamento", "valor": valor})

def ler_saft(entrada, desde, ate, pag_ate, toques):
    """Faturas (FT, FS, FR), notas de crédito (NC) e de débito (ND), sem as anuladas; recibos até pag_ate."""
    docs = []
    for f in sorted(glob.glob(os.path.join(entrada, "saft", "*.xml"))):
        try: root = ET.parse(f).getroot()
        except ET.ParseError as e:
            print(f"aviso: SAF-T inválido {f}: {e}", file=sys.stderr); continue
        ns = {"n": root.tag.split("}")[0].strip("{")} if "}" in root.tag else {}
        p = (lambda t: f"n:{t}") if ns else (lambda t: t)
        q = lambda caminho: "/".join(p(x) for x in caminho.split("/"))
        clientes = {c.findtext(p("CustomerID"), namespaces=ns): c.findtext(p("CompanyName"), namespaces=ns)
                    for c in root.findall(".//" + q("MasterFiles/Customer"), ns)}
        for inv in root.findall(".//" + q("SalesInvoices/Invoice"), ns):
            if (inv.findtext(q("DocumentStatus/InvoiceStatus"), namespaces=ns) or "N") == "A": continue
            tipo = inv.findtext(p("InvoiceType"), namespaces=ns) or ""
            d = data_iso(inv.findtext(p("InvoiceDate"), namespaces=ns))
            cid = inv.findtext(p("CustomerID"), namespaces=ns)
            nome = clientes.get(cid, cid)
            tot = float(inv.findtext(q("DocumentTotals/NetTotal"), namespaces=ns) or 0)
            no = inv.findtext(p("InvoiceNo"), namespaces=ns)
            docs.append({"no": no, "tipo": tipo, "data": d, "cliente": nome, "liquido": tot})
            if tipo in ("FT", "FS", "FR") and no_periodo(d, desde, ate):
                toques.append({"fonte": "saft", "data": d, "cliente": nome, "detalhe": no, "etapa": "fatura", "valor": tot})
                if tipo == "FR":  # fatura-recibo: paga no ato
                    toques.append({"fonte": "saft", "data": d, "cliente": nome, "detalhe": no, "etapa": "pagamento"})
        # só os recibos da secção Payments (não os <Payment> dentro dos totais de cada fatura), sem os anulados
        for pay in root.findall(".//" + q("SourceDocuments/Payments/Payment"), ns):
            if (pay.findtext(q("DocumentStatus/PaymentStatus"), namespaces=ns) or "N") == "A": continue
            d = data_iso(pay.findtext(p("TransactionDate"), namespaces=ns))
            if not no_periodo(d, desde, pag_ate): continue
            cid = pay.findtext(p("CustomerID"), namespaces=ns)
            toques.append({"fonte": "saft", "data": d, "cliente": clientes.get(cid, cid), "detalhe": pay.findtext(p("PaymentRefNo"), namespaces=ns) or "recibo", "etapa": "pagamento"})
    return docs

def ler_registo(entrada, desde, ate, toques):
    for r in ler_csv(os.path.join(entrada, "pedidos.csv")):
        cli = r.get("cliente") or ""
        for etapa, col in (("pedido", "data_pedido"), ("visita", "data_visita"), ("orcamento", "data_envio"), ("adjudicacao", "data_aceite")):
            d = data_iso(r.get(col))
            if d and cli: toques.append({"fonte": "registo", "data": d, "cliente": cli, "detalhe": r.get("num") or r.get("id"), "etapa": etapa, "valor": _num(r.get("valor"))})
        if r.get("estado") == "perdido" and data_iso(r.get("data_fecho")):
            toques.append({"fonte": "registo", "data": data_iso(r.get("data_fecho")), "cliente": cli, "detalhe": r.get("motivo") or "perdido", "etapa": "perdido"})

# ---------- agrupar por cliente e dividir em pedidos ----------
def agrupar(toques):
    grupos = []  # cada grupo: {"nomes": set, "toques": []}
    for t in sorted(toques, key=lambda x: x["data"] or ""):
        if not t.get("cliente") or not t.get("data"): continue
        alvo = next((g for g in grupos if any(parece_mesmo(t["cliente"], n) for n in g["nomes"])), None)
        if alvo is None:
            alvo = {"nomes": set(), "toques": []}; grupos.append(alvo)
        alvo["nomes"].add(t["cliente"]); alvo["toques"].append(t)
    return grupos

def dividir(g, intervalo):
    """Um cliente pode ter vários pedidos no período. Abre-se outro quando chega um pedido novo depois de o anterior
    estar faturado ou perdido, ou quando passam mais de «intervalo» dias sem nada. Os pagamentos vão para o último
    pedido faturado e ainda por pagar, mesmo que cheguem tarde."""
    pedidos = []
    for t in g["toques"]:
        atual = pedidos[-1] if pedidos else None
        if t["etapa"] == "pagamento":
            alvo = next((x for x in reversed(pedidos) if any(y["etapa"] == "fatura" for y in x["toques"])
                         and not any(y["etapa"] == "pagamento" for y in x["toques"])), atual)
            if alvo is None:
                alvo = {"nomes": g["nomes"], "toques": []}; pedidos.append(alvo)
            alvo["toques"].append(t); continue
        if atual is not None:
            fechado = any(y["etapa"] in ("fatura", "perdido") for y in atual["toques"])
            parado = dias(atual["toques"][-1]["data"], t["data"]) > intervalo
            if (t["etapa"] == "pedido" and fechado) or parado: atual = None
        if atual is None:
            atual = {"nomes": g["nomes"], "toques": []}; pedidos.append(atual)
        atual["toques"].append(t)
    return pedidos

def percurso(g):
    linha = {"cliente": sorted(g["nomes"], key=len)[-1], "fontes": ",".join(sorted({t["fonte"] for t in g["toques"]}))}
    for e in ETAPAS: linha[e] = None
    linha["perdido"] = None; linha["valor_orc"] = None; linha["valor_fat"] = 0.0; linha["inferido"] = []
    por_etapa = {}
    for t in g["toques"]:
        e = t["etapa"]
        if e == "perdido": linha["perdido"] = linha["perdido"] or t["data"]; continue
        if e not in ETAPAS: continue
        por_etapa.setdefault(e, []).append((t["data"], t["fonte"]))
        if linha[e] is None or t["data"] < linha[e]: linha[e] = t["data"]
        if e == "orcamento" and t.get("valor"): linha["valor_orc"] = t["valor"]
        if e == "fatura" and t.get("valor"): linha["valor_fat"] += t["valor"]
    # Duas fontes com datas diferentes para a mesma etapa: guardam-se as duas, não se escolhe.
    conflitos = []
    for e, lst in por_etapa.items():
        primeiras = {}
        for d, f in sorted(lst):
            primeiras.setdefault(f, d)
        if len(set(primeiras.values())) > 1:
            conflitos.append(f"{e}: " + " vs ".join(f"{d} ({f})" for f, d in sorted(primeiras.items(), key=lambda x: x[1])))
    if linha["perdido"] and linha["fatura"]:
        conflitos.append(f"perdido a {linha['perdido']} e faturado a {linha['fatura']}")
    if linha["fatura"] and not linha["adjudicacao"]: linha["adjudicacao"] = linha["obra"] or linha["fatura"]; linha["inferido"].append("adjudicacao")
    if linha["fatura"] and not linha["obra"]: linha["obra"] = linha["fatura"]; linha["inferido"].append("obra")
    linha["conflitos"] = "; ".join(conflitos)
    linha["inferido"] = ",".join(linha["inferido"])
    return linha

# ---------- relatórios ----------
def escrever(saida, linhas, docs, chamadas, meta, desde, ate, sem_data):
    os.makedirs(saida, exist_ok=True)
    cols = ["cliente", "fontes"] + ETAPAS + ["perdido", "valor_orc", "valor_fat", "inferido", "conflitos"]
    with open(os.path.join(saida, "pedidos_reconstruidos.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols); w.writeheader()
        for l in linhas: w.writerow({k: l.get(k) for k in cols})

    com_pedido = [l for l in linhas if l["pedido"] and no_periodo(l["pedido"], desde, ate)]
    n = {e: sum(1 for l in com_pedido if l[e]) for e in ETAPAS}
    perdidos = sum(1 for l in com_pedido if l["perdido"] and not l["fatura"])
    abertos = sum(1 for l in com_pedido if not l["fatura"] and not l["perdido"])
    def pct(a, b): return f"{(100*a/b):.0f}%" if b else "n/a"
    med = {f"{a}→{b}": mediana([dias(l[a], l[b]) for l in com_pedido if l[a] and l[b]]) for a, b in zip(ETAPAS, ETAPAS[1:])}
    eur_orc = sum(l["valor_orc"] or 0 for l in com_pedido)
    eur_ganho = sum(l["valor_orc"] or 0 for l in com_pedido if l["adjudicacao"] and not (l["perdido"] and not l["fatura"]))
    eur_perdido = sum(l["valor_orc"] or 0 for l in com_pedido if l["perdido"] and not l["fatura"])
    eur_fat = sum(l["valor_fat"] for l in com_pedido)
    eur_por_cobrar = sum(l["valor_fat"] for l in com_pedido if l["fatura"] and not l["pagamento"])
    sem_pedido = [l for l in linhas if l["fatura"] and not l["pedido"]]
    clientes = len({l["cliente"] for l in linhas})

    md = [f"# Funil {desde} a {ate}\n",
          f"Fontes lidas: {meta.get('emails', 0)} emails de clientes ({meta.get('emails_excluidos', 0)} excluídos: fornecedores, MDM, avisos), "
          f"{len(chamadas)} chamadas, {len(docs)} documentos SAF-T. {len(linhas)} pedidos de {clientes} clientes.\n",
          "| Etapa | N | % da anterior | % do pedido |", "|---|---|---|---|"]
    ant = None
    for e in ETAPAS:
        md.append(f"| {e} | {n[e]} | {pct(n[e], ant) if ant is not None else ''} | {pct(n[e], n['pedido'])} |"); ant = n[e]
    md += [f"| perdido | {perdidos} | | {pct(perdidos, n['pedido'])} |", f"| ainda aberto | {abertos} | | {pct(abertos, n['pedido'])} |", "",
           "Adjudicação e obra sem registo próprio são inferidas da fatura (coluna inferido de pedidos_reconstruidos.csv): "
           f"{sum(1 for l in com_pedido if 'adjudicacao' in l['inferido'])} adjudicações inferidas.", "",
           "## Dias medianos entre etapas", "| De → para | Mediana (dias) | Com dados |", "|---|---|---|"]
    for k, v in med.items():
        a, b = k.split("→"); c = sum(1 for l in com_pedido if l[a] and l[b])
        md.append(f"| {k} | {num1(v)} | {c} |")
    md += ["", "## Euros (s/IVA, só onde há valor lido)",
           f"- Orçamentado: {eur(eur_orc)} ({sum(1 for l in com_pedido if l['valor_orc'])} orçamentos com valor)",
           f"- Ganho (orçamento com adjudicação): {eur(eur_ganho)}", f"- Perdido (orçamento marcado perdido, sem fatura): {eur(eur_perdido)}",
           f"- Faturado a estes pedidos: {eur(eur_fat)}", f"- Faturado e sem pagamento registado: {eur(eur_por_cobrar)}", "",
           f"## Faturas sem pedido encontrado: {len(sem_pedido)}",
           "Pedidos faturados no período sem email, WhatsApp, agenda ou ORC que explique de onde veio o trabalho (clientes recorrentes ou pedidos por telefone sem registo). Ver pedidos_reconstruidos.csv, coluna fontes = saft.", "",
           "## Conflitos entre fontes",
           f"{sum(1 for l in linhas if l['conflitos'])} pedidos com datas diferentes para a mesma etapa em fontes diferentes, ou perdidos e faturados. Estão na coluna conflitos; não se escolheu nenhuma.", "",
           "## Orçamentos sem data", f"{len(sem_data)} ficheiros ORC sem data na folha nem no nome: ficaram fora do funil. Lista em orc_sem_data.csv."]
    gravar(os.path.join(saida, "funil.md"), "\n".join(md) + "\n")

    # reconciliação SAF-T: o líquido é faturas menos notas de crédito (mais notas de débito), sem anuladas
    por_mes = {}
    for d in docs:
        if not no_periodo(d["data"], desde, ate): continue
        m = por_mes.setdefault(d["data"][:7], {"nf": 0, "vf": 0.0, "nc": 0, "vc": 0.0, "nd": 0, "vd": 0.0})
        if d["tipo"] in ("FT", "FS", "FR"): m["nf"] += 1; m["vf"] += d["liquido"]
        elif d["tipo"] == "NC": m["nc"] += 1; m["vc"] += d["liquido"]
        elif d["tipo"] == "ND": m["nd"] += 1; m["vd"] += d["liquido"]
    fat_periodo = [d for d in docs if d["tipo"] in ("FT", "FS", "FR") and no_periodo(d["data"], desde, ate)]
    casadas = sum(1 for f in fat_periodo if any(f["cliente"] and parece_mesmo(f["cliente"], l["cliente"]) and l["pedido"] for l in linhas))
    rc = [f"# Reconciliação com o SAF-T {desde} a {ate}\n",
          "| Mês | Faturas | Valor faturas | Notas de crédito | Valor NC | Líquido do mês |", "|---|---|---|---|---|---|"]
    tf = tv = tn = tc = tl = 0
    for m in sorted(por_mes):
        x = por_mes[m]; liq = x["vf"] - x["vc"] + x["vd"]
        rc.append(f"| {m} | {x['nf']} | {eur(x['vf'])} | {x['nc']} | {eur(x['vc'])} | {eur(liq)} |")
        tf += x["nf"]; tv += x["vf"]; tn += x["nc"]; tc += x["vc"]; tl += liq
    rc += [f"| **Total** | **{tf}** | **{eur(tv)}** | **{tn}** | **{eur(tc)}** | **{eur(tl)}** |", "",
           "Anuladas não contam. O «Líquido do mês» é o que se compara com o contabilista (faturas menos notas de crédito, mais notas de débito).", "",
           f"Faturas com pedido reconstruído a montante: {casadas} de {tf} ({pct(casadas, tf)}).",
           f"Faturado a pedidos reconstruídos (funil.md): {eur(eur_fat)}. Diferença para o valor das faturas: {eur(tv - eur_fat)}.",
           "", "A diferença é trabalho faturado cujo pedido não deixou rasto em nenhum canal exportado. Se for grande, o problema é de registo, não de vendas."]
    gravar(os.path.join(saida, "reconciliacao.md"), "\n".join(rc) + "\n")

    # chamadas
    cm = {}
    for c in chamadas:
        k = cm.setdefault((c["mes"], c["linha"]), {"recebidas": 0, "perdidas": 0})
        if c["recebida"]: k["recebidas"] += 1
        if c["perdida"]: k["perdidas"] += 1
    ch = ["# Chamadas\n", "| Mês | Linha | Recebidas | Perdidas | % perdidas |", "|---|---|---|---|---|"]
    for (m, l), v in sorted(cm.items()): ch.append(f"| {m} | {l} | {v['recebidas']} | {v['perdidas']} | {pct(v['perdidas'], v['recebidas'])} |")
    if not cm: ch.append("| sem ficheiros chamadas-*.csv | | | | |")
    ch += ["", "Chamadas que viraram pedido: não é possível saber a partir da fatura do operador. É o que a folha de contagem de um dia (02-observacao) mede."]
    gravar(os.path.join(saida, "chamadas.md"), "\n".join(ch) + "\n")

    # não correspondidos, remetentes e orçamentos sem data (têm nomes: ficam em saida/, fora do git)
    with open(os.path.join(saida, "nao_correspondidos.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["cliente", "fontes", "etapas_encontradas"])
        for l in linhas:
            if l["fontes"] in ("email", "whatsapp") or (l["fatura"] and not l["pedido"]):
                w.writerow([l["cliente"], l["fontes"], ",".join(e for e in ETAPAS if l[e])])
    with open(os.path.join(saida, "remetentes.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["dominio", "emails"])
        for dom, k in meta.get("dominios", Counter()).most_common(): w.writerow([dom, k])
    with open(os.path.join(saida, "orc_sem_data.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["ficheiro", "num"]); w.writeheader(); w.writerows(sem_data)

# ---------- demo ----------
def gerar_demo(entrada):
    import random; random.seed(7)
    for p in ("whatsapp/Zeta Particular", "orc", "saft"): os.makedirs(os.path.join(entrada, p), exist_ok=True)
    clientes = ["Alfa Construções Lda", "Beta Escola", "Gama Condomínio", "Delta Banco", "Épsilon Indústria", "Zeta Particular", "Eta Clínica", "Teta Restaurante"]
    mb = mailbox.mbox(os.path.join(entrada, "gmail.mbox"))
    for i, c in enumerate(clientes[:6]):
        d = dt.datetime(2026, 7, 1) + dt.timedelta(days=random.randint(0, 80))
        m = mailbox.mboxMessage(); m["From"] = f"{c} <c{i}@exemplo.pt>"; m["Subject"] = "Pedido de orcamento AC"; m["Date"] = d.strftime("%a, %d %b %Y 10:00:00 +0100"); m.set_payload("Ola", charset="utf-8"); mb.add(m)
    m = mailbox.mboxMessage(); m["From"] = "Fornecedor Exemplo <vendas@fornecedor-exemplo.pt>"; m["Subject"] = "Tabela de preços"; m["Date"] = "Mon, 13 Jul 2026 10:00:00 +0100"; m.set_payload("x"); mb.add(m)
    mb.flush()
    gravar(os.path.join(entrada, "remetentes-excluir.txt"), "# fornecedores e outros que não são clientes\nfornecedor-exemplo.pt\n")
    gravar(os.path.join(entrada, "whatsapp", "Zeta Particular", "_chat.txt"), "[12/07/2026, 09:15:22] Zeta Particular: Bom dia, preciso de um AC\n[12/07/2026, 09:20:00] MDM: Bom dia, quando podemos ir?\n")
    with open(os.path.join(entrada, "chamadas-movel.csv"), "w", newline="") as f:
        w = csv.writer(f); w.writerow(["data", "hora", "numero", "duracao_s", "tipo"])
        for _ in range(120):
            d = dt.date(2026, 7, 1) + dt.timedelta(days=random.randint(0, 91)); t = random.choice(["recebida", "recebida", "recebida", "perdida", "efetuada"])
            w.writerow([d.isoformat(), "10:00", "9xxxxxxxx", 0 if t == "perdida" else random.randint(20, 400), t])
    with open(os.path.join(entrada, "agenda.csv"), "w", newline="") as f:
        w = csv.writer(f); w.writerow(["Data", "Cliente", "Técnico"])
        for c in clientes[:4]: w.writerow(["2026-07-" + str(random.randint(10, 28)).zfill(2), c, "Carlos"])
    try:
        import openpyxl
        for i, c in enumerate(clientes[:3]):
            wb = openpyxl.Workbook(); ws = wb.active; ws["A3"] = "Data"; ws["B3"] = dt.datetime(2026, 8, 5 + i); ws["B20"] = f"Valor Total: {1200+i*300},00 € + IVA"
            wb.save(os.path.join(entrada, "orc", f"ORC00{60+i}_2026 - {c}.xlsx"))
    except ImportError: pass
    ns = "urn:OECD:StandardAuditFile-Tax:PT_1.04_01"
    xml = [f'<AuditFile xmlns="{ns}"><MasterFiles>']
    for i, c in enumerate([clientes[0], clientes[2], clientes[7]]): xml.append(f"<Customer><CustomerID>C{i}</CustomerID><CompanyName>{c}</CompanyName></Customer>")
    xml.append("</MasterFiles><SourceDocuments><SalesInvoices>")
    for i, (no, tipo, d, tot, cid) in enumerate([("FT 2026/40", "FT", "2026-08-20", "1200.00", "C0"), ("FR 2026/41", "FR", "2026-09-02", "1500.00", "C1"),
                                                 ("FT 2026/42", "FT", "2026-09-15", "320.00", "C2"), ("NC 2026/5", "NC", "2026-09-20", "120.00", "C2"),
                                                 ("FT 2026/43", "FT", "2026-09-21", "999.00", "C2")]):
        st = "A" if no == "FT 2026/43" else "N"
        xml.append(f"<Invoice><InvoiceNo>{no}</InvoiceNo><DocumentStatus><InvoiceStatus>{st}</InvoiceStatus></DocumentStatus><InvoiceDate>{d}</InvoiceDate><InvoiceType>{tipo}</InvoiceType><CustomerID>{cid}</CustomerID><DocumentTotals><NetTotal>{tot}</NetTotal></DocumentTotals></Invoice>")
    xml.append("</SalesInvoices><Payments><Payment><PaymentRefNo>RC 1</PaymentRefNo><TransactionDate>2026-10-12</TransactionDate><DocumentStatus><PaymentStatus>N</PaymentStatus></DocumentStatus><CustomerID>C0</CustomerID></Payment></Payments></SourceDocuments></AuditFile>")
    gravar(os.path.join(entrada, "saft", "DadosFat_2026-09.xml"), "".join(xml))
    with open(os.path.join(entrada, "pedidos.csv"), "w", newline="") as f:
        w = csv.writer(f); w.writerow(["id", "num", "cliente", "tipo", "origem", "estado", "prioridade", "dono", "data_pedido", "data_visita", "data_envio", "data_aceite", "data_fecho", "motivo", "valor", "por_confirmar", "criado"])
        w.writerow(["x1", "ORC0061_2026", "Beta Escola", "instalacao", "email", "perdido", "P2", "B", "2026-07-15", "", "2026-08-06", "", "2026-09-01", "Preço", "1500.5", "false", "2026-10-16"])
        w.writerow(["x2", "ORC0060_2026", "Alfa Construções", "instalacao", "", "em_execucao", "P2", "B", "2026-07-03", "", "", "2026-08-10", "", "", "1200", "false", "2026-10-16"])

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--entrada", default="entrada"); ap.add_argument("--saida", default="saida")
    ap.add_argument("--desde", default="2026-07-01"); ap.add_argument("--ate", default="2026-09-30")
    ap.add_argument("--pagamentos-ate", default=None, help="até quando contam recibos de faturas do período (por omissão, 120 dias depois de --ate)")
    ap.add_argument("--intervalo", type=int, default=60, help="dias sem nada que separam dois pedidos do mesmo cliente")
    ap.add_argument("--eu", default="MDM", help="nomes da MDM nas conversas de WhatsApp, separados por vírgula")
    ap.add_argument("--demo", action="store_true", help="gera dados fictícios em entrada-demo/ e corre")
    a = ap.parse_args(argv)
    if a.demo:
        a.entrada, a.saida = "entrada-demo", "saida-demo"
        if os.path.exists(a.entrada):
            import shutil; shutil.rmtree(a.entrada)
        gerar_demo(a.entrada)
    pag_ate = a.pagamentos_ate or (dt.date.fromisoformat(a.ate) + dt.timedelta(days=120)).isoformat()
    toques, meta, sem_data = [], {}, []
    ler_emails(a.entrada, a.desde, a.ate, toques, meta, ler_excluir(a.entrada))
    ler_whatsapp(a.entrada, a.desde, a.ate, toques, [x.strip() for x in a.eu.split(",") if x.strip()])
    chamadas = ler_chamadas(a.entrada, a.desde, a.ate)
    ler_agenda(a.entrada, a.desde, a.ate, toques)
    ler_orc(a.entrada, a.desde, a.ate, toques, sem_data)
    docs = ler_saft(a.entrada, a.desde, a.ate, pag_ate, toques)
    ler_registo(a.entrada, a.desde, a.ate, toques)
    linhas = [percurso(p) for g in agrupar(toques) for p in dividir(g, a.intervalo)]
    escrever(a.saida, linhas, docs, chamadas, meta, a.desde, a.ate, sem_data)
    print(f"{len(toques)} toques, {len(linhas)} pedidos, {len(docs)} documentos SAF-T, {len(chamadas)} chamadas → {a.saida}/")
    with open(os.path.join(a.saida, "funil.md"), encoding="utf-8") as f: print(f.read())

if __name__ == "__main__":
    main()
