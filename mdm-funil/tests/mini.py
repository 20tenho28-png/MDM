"""Pequenos projetos de teste, construídos à mão com poucos ficheiros em dados/."""

from __future__ import annotations

import itertools
import mailbox
from datetime import datetime
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

from openpyxl import Workbook

from funil.datas import LX

MDM = '"MDM Climatização" <mdmassistencia@gmail.com>'

_CONFIG = """
[geral]
data_referencia = "{hoje}"
[empresa]
emails = ["mdmassistencia@gmail.com"]
telefones = ["+351219990000"]
nomes_whatsapp = [{nomes}]
"""

_ids = itertools.count(1)


def projeto(raiz: Path, hoje: str = "2026-10-16", nomes=("MDM Climatização",), extra: str = "") -> Path:
    for sub in ("email", "whatsapp", "chamadas", "orc", "faturas", "site"):
        (raiz / "dados" / sub).mkdir(parents=True, exist_ok=True)
    lista = ", ".join(f'"{n}"' for n in nomes)
    (raiz / "config.toml").write_text(_CONFIG.format(hoje=hoje, nomes=lista) + extra, encoding="utf-8")
    return raiz


def t(*a) -> datetime:
    return datetime(*a, tzinfo=LX)


def email(de: str, para: str, quando: datetime, assunto: str, corpo: str, anexo: str | None = None,
          **cabecalhos) -> EmailMessage:
    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = de, para, assunto
    m["Date"] = format_datetime(quando)
    m["Message-ID"] = f"<teste{next(_ids)}@exemplo.pt>"
    for k, v in cabecalhos.items():
        m[k.replace("_", "-")] = v
    m.set_content(corpo)
    if anexo:
        m.add_attachment(b"%PDF-1.4", maintype="application", subtype="pdf", filename=anexo)
    return m


def caixa(raiz: Path, mensagens, nome: str = "caixa.mbox") -> None:
    c = mailbox.mbox(str(raiz / "dados" / "email" / nome))
    for m in mensagens:
        c.add(m)
    c.flush()
    c.close()


def orcamento(raiz: Path, numero: str, cliente: str, data=None, valor: str = "1.000,00 € + IVA",
              contacto: str | None = None, criado: datetime | None = None) -> Path:
    """Ficheiro ORC parecido com o da MDM: cliente, contacto, data e 'Valor Total'."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Orçamento"
    ws["A4"], ws["B4"] = "Cliente:", cliente
    if contacto:
        ws["A5"], ws["B5"] = "Contacto:", contacto
    if data is not None:
        ws["A6"], ws["B6"] = "Data:", data
    ws["A12"] = f"Valor Total: {valor}"
    if criado is not None:
        wb.properties.created = criado
    caminho = raiz / "dados" / "orc" / f"{numero} {cliente}.xlsx"
    wb.save(caminho)
    return caminho


def whatsapp(raiz: Path, contacto: str, linhas: list[str]) -> Path:
    """Exportação de Android: 'Conversa do WhatsApp com <contacto>.txt'."""
    caminho = raiz / "dados" / "whatsapp" / f"Conversa do WhatsApp com {contacto}.txt"
    caminho.write_text("\n".join(linhas) + "\n", encoding="utf-8")
    return caminho
