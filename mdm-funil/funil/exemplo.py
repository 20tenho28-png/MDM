"""Mundo de exemplo: dados inventados, com as respostas certas guardadas, para testar o funil e ver o relatório.

Nada aqui é da MDM. Os clientes chamam-se "<Nome> Exemplo <número>" e os números de telefone são fictícios.
"""

from __future__ import annotations

import csv
import mailbox
import math
import random
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

from .datas import LX, Calendario
from .texto import sem_acentos

MDM_EMAIL = "mdmassistencia@gmail.com"
MDM_TEL = "+351219990000"
MDM_ASSINATURA = "Cumprimentos,\nMDM Climatização\nTel. 219 990 000\nmdmassistencia@gmail.com"
MDM_WA = "MDM Climatização"
FORNECEDOR_TEL = "+351210001111"
FORNECEDOR_EMAIL = "comercial@distribuidor-exemplo.pt"
# Armadilhas reais: o técnico recebe os números ORC por email, o arquiteto vai em cópia, há férias com resposta
# automática e o WhatsApp Business manda uma saudação. Nada disto pode mudar os números do funil.
TECNICO = '"Carlos Técnico" <carlos.tecnico.mdm@gmail.com>'
ARQUITETO = '"Arq. Sofia Almeida" <sofia@almeida-arquitetos.pt>'
FERIAS = (date(2026, 8, 10), date(2026, 8, 21))
SAUDACAO_WA = "Olá! Obrigado por contactar a MDM Climatização. Respondemos com a maior brevidade possível."

PRIMEIROS = ["Ana", "Bruno", "Carla", "Diogo", "Eva", "Filipe", "Gabriela", "Hugo", "Inês", "João", "Luísa",
             "Marco", "Nuno", "Olga", "Paulo", "Rita", "Sérgio", "Teresa", "Vasco", "Zita", "Tiago", "Sofia"]
LOCAIS = ["Lisboa", "Almada", "Seixal", "Amadora", "Odivelas", "Loures", "Oeiras", "Barreiro", "Montijo", "Sintra"]

TIPOS = [("instalacao", 0.5), ("avaria", 0.2), ("manutencao", 0.2), ("obra", 0.1)]
CANAIS = [("email", 0.40), ("whatsapp", 0.25), ("chamada", 0.20), ("site", 0.15)]
# Probabilidade de ganhar com resposta imediata e quanto "aguenta" a espera (horas úteis): p = base x e^(-h/tau).
BASE = {"avaria": 0.75, "instalacao": 0.45, "manutencao": 0.55, "obra": 0.30}
TAU = {"avaria": 8.0, "instalacao": 40.0, "manutencao": 90.0, "obra": 150.0}
P_ORCAMENTO = {"instalacao": 0.85, "obra": 0.9, "manutencao": 0.6, "avaria": 0.3}

TEXTOS = {
    "instalacao": [
        "Boa tarde, gostaria de pedir um orçamento para a instalação de {u} {equip}. Moro em {local}.",
        "Olá! Queria saber quanto fica instalar {u} {equip} (sala e quartos). Obrigado.",
        "Bom dia, estou a pensar comprar ar condicionado para {u} divisões. Podem fazer orçamento com instalação?",
    ],
    "avaria": [
        "Bom dia, o ar condicionado deixou de funcionar e está a pingar. Podem vir ver?{urg}",
        "Olá, o nosso AC não arrefece e dá código de erro E5.{urg}",
        "Boa tarde, a máquina exterior faz muito barulho e parou de aquecer.{urg}",
    ],
    "manutencao": [
        "Bom dia, gostaria de marcar a manutenção e limpeza de {u} {equip}.",
        "Olá, precisamos da revisão anual de {u} aparelhos.",
    ],
    "obra": [
        "Boa tarde, somos uma construtora e precisamos de orçamento de climatização para uma obra com {u} frações em {local}.",
        "Bom dia, enviamos o caderno de encargos da empreitada ({u} frações). Aguardamos a vossa proposta.",
    ],
}
FRASE_EMPRESA = ["É para o nosso escritório.", "É para a nossa loja.", "Somos um restaurante.", "É para a clínica."]
URGENCIA = [" É urgente, estamos com muito calor.", " Precisamos com urgência."]
ACEITES = ["Bom dia, aceito o orçamento. Quando podem vir instalar?", "Ok, pode avançar. Quando podem vir?",
           "Boa tarde, vamos avançar com a vossa proposta."]
RECUSAS = {
    "foi com outra empresa": "Obrigado pelo orçamento, mas já adjudicámos a outra empresa.",
    "preço": "Obrigado, mas ficou um pouco caro para nós.",
    "desistiu ou adiou": "Por agora vamos desistir, obrigado pela atenção.",
}
# Clientes que cobram quando a resposta ou o orçamento demoram.
COBRA_RESPOSTA = ["Bom dia, então, já viram o meu pedido?", "Olá, alguma novidade?",
                  "Bom dia, ainda estou à espera de resposta."]
COBRA_ORCAMENTO = ["Bom dia, já têm o orçamento?", "Olá, alguma novidade sobre o orçamento?",
                   "Boa tarde, ainda não recebi o orçamento."]
ASSUNTOS = {"instalacao": "Pedido de orçamento", "avaria": "Avaria ar condicionado",
            "manutencao": "Manutenção AC", "obra": "Orçamento climatização obra"}


@dataclass
class Cliente:
    id: int
    nome: str
    tel: str
    email: str | None
    empresa: bool
    nif: str
    nome_no_telefone: str = ""
    formato_wa: str = "android"
    ultimo_contacto: datetime | None = None


@dataclass
class Verdade:
    cliente: Cliente
    canal: str
    tipo: str
    unidades: int
    urgente: bool
    t_contacto: datetime
    texto: str
    atendida: bool = False
    t_resposta: datetime | None = None
    h_resposta: float | None = None
    orc: str = ""
    t_envio: datetime | None = None
    via_orc: str = ""
    valor: float | None = None
    ganho: bool = False
    t_aceite: datetime | None = None
    aceite_por_texto: bool = False
    t_recusa: datetime | None = None
    motivo: str = ""
    t_fatura: datetime | None = None
    fio: str = ""
    cobrancas: list = field(default_factory=list)   # (momento, texto) das mensagens a cobrar resposta


def _poisson(rng: random.Random, lam: float) -> int:
    limite, k, p = math.exp(-lam), 0, 1.0
    while True:
        p *= rng.random()
        if p <= limite:
            return k
        k += 1


def _escolher(rng: random.Random, pesos):
    r, acumulado = rng.random(), 0.0
    for valor, peso in pesos:
        acumulado += peso
        if r <= acumulado:
            return valor
    return pesos[-1][0]


def _minuto(t: datetime) -> datetime:
    return t.replace(second=0, microsecond=0)


def _tel_txt(tel: str) -> str:
    d = tel[4:]
    return f"{d[:3]} {d[3:6]} {d[6:]}"


def _pt(v: float) -> str:
    return f"{v:,.2f}".replace(",", " ").replace(".", ",").replace(" ", ".")


class Mundo:
    def __init__(self, semente: int = 7, inicio: date = date(2025, 10, 6), fim: date = date(2026, 10, 9),
                 hoje: date = date(2026, 10, 16), por_semana: float = 6.0):
        self.rng = random.Random(semente)
        self.ruido = random.Random(semente + 1000)   # só para as armadilhas, para não mudar o resto do mundo
        self.cobra = random.Random(semente + 2000)   # só para as cobranças dos clientes, pela mesma razão
        self.visitas = random.Random(semente + 3000)  # só para a agenda do ERP, pela mesma razão
        self.cal = Calendario(feriados_extra=["06-13"])
        self.inicio, self.fim, self.hoje, self.por_semana = inicio, fim, hoje, por_semana
        self.clientes: list[Cliente] = []
        self.pedidos: list[Verdade] = []

    # ------------------------------------------------------------------ simulação

    def _cliente(self, quando: datetime) -> Cliente:
        rng = self.rng
        antigos = [c for c in self.clientes if c.ultimo_contacto and (quando - c.ultimo_contacto).days > 200]
        if antigos and rng.random() < 0.05:
            return rng.choice(antigos)
        i = len(self.clientes) + 1
        primeiro = rng.choice(PRIMEIROS)
        empresa = rng.random() < 0.22
        tel = f"+3519100{i:05d}"
        if empresa:
            email = f"geral@empresa-exemplo-{i:03d}.pt"
        else:
            email = f"{sem_acentos(primeiro).lower()}.exemplo{i:03d}@gmail.com" if rng.random() < 0.8 else None
        c = Cliente(i, f"{primeiro} Exemplo {i:03d}", tel, email, empresa, f"{5 if empresa else 2}{i:08d}",
                    formato_wa="iphone" if rng.random() < 0.3 else "android")
        c.nome_no_telefone = c.nome if rng.random() < 0.3 else ""
        self.clientes.append(c)
        return c

    def _instante_contacto(self, d: date, canal: str) -> datetime:
        rng = self.rng
        if canal == "chamada":
            h = rng.uniform(9, 18) if self.cal.e_util(d) else rng.uniform(10, 13)
        elif canal == "whatsapp":
            h = rng.uniform(8, 22)
        else:
            h = rng.uniform(7, 23)
        return _minuto(datetime.combine(d, time(0), tzinfo=LX) + timedelta(hours=h))

    def _texto(self, tipo: str, u: int, empresa: bool, urgente: bool) -> str:
        rng = self.rng
        equip = "ar condicionado" if u == 1 else "aparelhos de ar condicionado"
        t = rng.choice(TEXTOS[tipo]).format(u=u, equip=equip, local=rng.choice(LOCAIS),
                                            urg=rng.choice(URGENCIA) if urgente else "")
        if empresa and tipo != "obra":
            t += " " + rng.choice(FRASE_EMPRESA)
        if tipo == "instalacao" and rng.random() < 0.1:
            t += " Precisamos até dia 15 porque é para uma mudança."
        return t

    def simular(self) -> None:
        rng = self.rng
        d = self.inicio
        while d <= self.fim:
            lam = self.por_semana / 7 * (1.25 if d.weekday() < 5 else 0.4)
            for _ in range(_poisson(rng, lam)):
                canal = _escolher(rng, CANAIS)
                t = self._instante_contacto(d, canal)
                c = self._cliente(t)
                if canal in ("email",) and not c.email:
                    canal = "whatsapp"
                tipo = _escolher(rng, TIPOS)
                if c.empresa and rng.random() < 0.3:
                    tipo = "obra"
                if tipo == "obra" and not c.empresa:
                    tipo = "instalacao"
                u = {"instalacao": rng.randint(1, 4), "manutencao": rng.randint(1, 5), "obra": rng.randint(6, 30),
                     "avaria": 1}[tipo]
                urgente = tipo == "avaria" and rng.random() < 0.5
                v = Verdade(c, canal, tipo, u, urgente, t, self._texto(tipo, u, c.empresa, urgente))
                v.fio = f"{rng.getrandbits(60):x}"
                c.ultimo_contacto = t
                self._responder(v)
                self.pedidos.append(v)
            d += timedelta(days=1)
        self._numerar_orcamentos()
        self._executar()
        self._cobrancas()

    def _responder(self, v: Verdade) -> None:
        rng = self.rng
        if v.canal == "chamada" and rng.random() < 0.6:
            v.atendida, v.t_resposta, v.h_resposta = True, v.t_contacto, 0.0
        elif rng.random() < 0.08:
            return
        else:
            h = min(150.0, max(0.1, math.exp(rng.gauss(math.log(6), 1.3))))
            v.t_resposta = _minuto(self.cal.somar_horas_uteis(v.t_contacto, h))
            v.h_resposta = self.cal.horas_uteis(v.t_contacto, v.t_resposta)
        h = v.h_resposta
        tau = TAU[v.tipo] * (0.6 if v.urgente else 1.0)
        p_ganho = BASE[v.tipo] * (1.08 if v.cliente.empresa else 1.0) * math.exp(-h / tau)
        if rng.random() < P_ORCAMENTO[v.tipo]:
            espera = rng.uniform(20, 80) if v.tipo == "obra" else rng.uniform(2, 50)
            v.t_envio = _minuto(self.cal.somar_horas_uteis(v.t_resposta, espera))
            v.via_orc = "email" if v.cliente.email and (v.canal in ("email", "site") or rng.random() < 0.3) else "whatsapp"
            v.valor = round({"instalacao": 950 + 780 * v.unidades + rng.gauss(0, 120),
                             "obra": 900 * v.unidades + rng.gauss(0, 800),
                             "manutencao": 50 + 70 * v.unidades,
                             "avaria": rng.uniform(140, 650)}[v.tipo], 2)
            if rng.random() < p_ganho:
                v.ganho = True
                v.t_aceite = _minuto(v.t_envio + timedelta(days=rng.uniform(1, 14)))
                v.aceite_por_texto = rng.random() < 0.75
            elif rng.random() < 0.45:
                v.motivo = rng.choice(list(RECUSAS))
                v.t_recusa = _minuto(v.t_envio + timedelta(days=rng.uniform(2, 25)))
        elif v.tipo in ("avaria", "manutencao") and rng.random() < p_ganho:
            v.ganho = True
            v.valor = round(rng.uniform(90, 450), 2)
            v.t_fatura = datetime.combine((v.t_resposta + timedelta(days=rng.randint(1, 6))).date(), time(18), tzinfo=LX)

    def _cobrancas(self) -> None:
        """Clientes de email e WhatsApp que cobram ("então, já viram?") quando a resposta ou o orçamento demoram."""
        rng = self.cobra
        fecho = datetime.combine(self.fim, time(18), tzinfo=LX)
        for v in self.pedidos:
            if v.canal not in ("email", "whatsapp"):
                continue
            fim = v.t_resposta or fecho
            if self.cal.horas_uteis(v.t_contacto, fim) > 12 and rng.random() < 0.5:
                quando = _minuto(self.cal.somar_horas_uteis(v.t_contacto, rng.uniform(9, 12)))
                if quando < fim:
                    v.cobrancas.append((quando, rng.choice(COBRA_RESPOSTA)))
            if v.t_resposta and v.t_envio and self.cal.dias_uteis(v.t_resposta, v.t_envio) > 4 and rng.random() < 0.4:
                quando = _minuto(self.cal.somar_horas_uteis(v.t_resposta, 9 * rng.uniform(3, 4)))
                if v.t_resposta < quando < v.t_envio:
                    v.cobrancas.append((quando, rng.choice(COBRA_ORCAMENTO)))

    def _numerar_orcamentos(self) -> None:
        por_ano: dict[int, int] = {}
        for v in sorted((v for v in self.pedidos if v.t_envio), key=lambda v: v.t_envio):
            ano = v.t_envio.year
            por_ano[ano] = por_ano.get(ano, 0) + 1
            v.orc = f"ORC{por_ano[ano]:04d}_{ano}"

    def _executar(self) -> None:
        """Trabalhos aceites: avarias e manutenções em poucos dias; instalações e obras numa fila de ~1 por semana."""
        rng = self.rng
        fila = sorted((v for v in self.pedidos if v.ganho and v.t_aceite), key=lambda v: v.t_aceite)
        cursor = None
        for v in fila:
            if v.tipo in ("avaria", "manutencao"):
                dia = (v.t_aceite + timedelta(days=rng.randint(1, 7))).date()
            else:
                inicio = (v.t_aceite + timedelta(days=4)).date()
                dia = max(inicio, cursor) if cursor else inicio
                cursor = dia + timedelta(days=rng.uniform(5, 9))
                cursor = date(cursor.year, cursor.month, cursor.day) if isinstance(cursor, datetime) else cursor
            v.t_fatura = datetime.combine(dia, time(18), tzinfo=LX)

    # ------------------------------------------------------------------ escrita

    def escrever(self, destino: Path) -> None:
        destino = Path(destino)
        dados = destino / "dados"
        for sub in ("email", "whatsapp", "chamadas", "orc", "faturas", "site", "agenda"):
            (dados / sub).mkdir(parents=True, exist_ok=True)
        visiveis = [v for v in self.pedidos]
        self._escrever_emails(dados / "email" / "Todo o correio, incluindo Spam e Lixo.mbox", visiveis)
        self._escrever_whatsapp(dados / "whatsapp", visiveis)
        self._escrever_chamadas(dados / "chamadas" / "chamadas-telemovel.xml", visiveis)
        self._escrever_site(dados / "site" / "pedidos_site.csv", visiveis)
        self._escrever_orcamentos(dados / "orc", visiveis)
        self._escrever_faturas(dados / "faturas" / "faturas.csv", visiveis)
        self._escrever_agenda(dados / "agenda" / "agenda.csv", visiveis)
        self._escrever_config(destino / "config.toml")
        self._escrever_verdade(destino / "verdade.csv")

    def _ate_hoje(self, t: datetime | None) -> bool:
        return t is not None and t.date() <= self.hoje

    def _msg(self, ts, de, para, assunto, corpo, fio, etiquetas, anexo=None, em_resposta=None, cabecalhos=None,
             rng=None):
        m = EmailMessage()
        m["From"], m["To"], m["Subject"] = de, para, assunto
        m["Date"] = format_datetime(ts)
        mid = f"<{(rng or self.rng).getrandbits(64):x}@exemplo.mdm>"
        m["Message-ID"] = mid
        if em_resposta:
            m["In-Reply-To"] = em_resposta
            m["References"] = em_resposta
        m["X-GM-THRID"] = fio
        m["X-Gmail-Labels"] = etiquetas
        for k, val in (cabecalhos or {}).items():
            m[k] = val
        m.set_content(corpo)
        if anexo:
            m.add_attachment(b"%PDF-1.4 exemplo", maintype="application", subtype="pdf", filename=anexo)
        return m, mid

    def _escrever_emails(self, caminho: Path, pedidos: list[Verdade]) -> None:
        rng = self.rng
        caixa = mailbox.mbox(str(caminho))
        caixa.lock()
        try:
            mdm = f'"MDM Climatização" <{MDM_EMAIL}>'
            for v in pedidos:
                c = v.cliente
                if not c.email:
                    continue
                cli = f'"{c.nome}" <{c.email}>'
                # Nas obras, o arquiteto da empresa vai em cópia nos emails da MDM.
                para = f"{cli}, {ARQUITETO}" if v.tipo == "obra" else cli
                assunto = ASSUNTOS[v.tipo]
                mid = None
                if v.canal == "email":
                    assinatura = f"\n\n{c.nome}" + (f"\nTel. {_tel_txt(c.tel)}" if rng.random() < 0.5 else "")
                    m, mid = self._msg(v.t_contacto, cli, mdm, assunto, v.texto + assinatura, v.fio,
                                       "Caixa de entrada,Importante")
                    caixa.add(m)
                    if FERIAS[0] <= v.t_contacto.date() <= FERIAS[1]:
                        m, _ = self._msg(v.t_contacto + timedelta(minutes=1), mdm, cli, "Re: " + assunto,
                                         "Obrigado pelo contacto. Estamos de férias até 21 de agosto e respondemos no "
                                         "regresso.", v.fio, "Enviados", em_resposta=mid,
                                         cabecalhos={"Auto-Submitted": "auto-replied"}, rng=self.ruido)
                        caixa.add(m)
                    for quando, texto in v.cobrancas:
                        if self._ate_hoje(quando):
                            m, _ = self._msg(quando, cli, mdm, "Re: " + assunto, f"{texto}\n\n{c.nome}", v.fio,
                                             "Caixa de entrada", em_resposta=mid, rng=self.cobra)
                            caixa.add(m)
                    if v.t_resposta and self._ate_hoje(v.t_resposta) and self.ruido.random() < 0.15:
                        meio = v.t_contacto + (v.t_resposta - v.t_contacto) / 2
                        m, _ = self._msg(meio, mdm, cli, "Re: " + assunto, "Bom dia, obrigado pelo contacto. Pod",
                                         v.fio, "Rascunhos", rng=self.ruido)
                        caixa.add(m)
                    if v.t_resposta and self._ate_hoje(v.t_resposta):
                        corpo = ("Bom dia,\n\nObrigado pelo contacto. Podemos agendar uma visita para avaliar. "
                                 f"Qual a sua disponibilidade?\n\n{MDM_ASSINATURA}\n\n"
                                 f"Em {v.t_contacto:%d/%m/%Y} às {v.t_contacto:%H:%M}, {c.nome} <{c.email}> escreveu:\n"
                                 + "\n".join("> " + l for l in v.texto.splitlines()))
                        m, _ = self._msg(v.t_resposta, mdm, para, "Re: " + assunto, corpo, v.fio, "Enviados", em_resposta=mid)
                        caixa.add(m)
                elif v.canal == "site" and v.t_resposta and self._ate_hoje(v.t_resposta) and v.via_orc != "chamada":
                    if c.id % 7 != 0:  # alguns clientes do site recebem chamada em vez de email
                        corpo = f"Bom dia,\n\nRecebemos o seu pedido através do site. Podemos agendar uma visita?\n\n{MDM_ASSINATURA}"
                        m, mid = self._msg(v.t_resposta, mdm, cli, "Pedido através do site", corpo, v.fio, "Enviados")
                        caixa.add(m)
                if v.orc and v.via_orc == "email" and self._ate_hoje(v.t_envio):
                    fio = f"{rng.getrandbits(60):x}"
                    corpo = f"Bom dia,\n\nSegue em anexo o orçamento {v.orc}, conforme combinado.\n\n{MDM_ASSINATURA}"
                    m, mid_orc = self._msg(v.t_envio, mdm, para, f"Orçamento {v.orc}", corpo, fio, "Enviados",
                                           anexo=f"{v.orc} {c.nome}.pdf")
                    caixa.add(m)
                    resposta = None
                    if v.ganho and v.aceite_por_texto and self._ate_hoje(v.t_aceite):
                        resposta = (v.t_aceite, rng.choice(ACEITES))
                    elif v.t_recusa and self._ate_hoje(v.t_recusa):
                        resposta = (v.t_recusa, RECUSAS[v.motivo])
                    if resposta:
                        citado = "\n".join("> " + l for l in corpo.splitlines())
                        texto = (f"{resposta[1]}\n\n{c.nome}\n\nEm {v.t_envio:%d/%m/%Y} às {v.t_envio:%H:%M}, "
                                 f"MDM Climatização <{MDM_EMAIL}> escreveu:\n{citado}")
                        m, _ = self._msg(resposta[0], cli, mdm, f"Re: Orçamento {v.orc}", texto, fio,
                                         "Caixa de entrada", em_resposta=mid_orc)
                        caixa.add(m)
            # O técnico recebe o número do orçamento para ir medir (e responde a citar o assunto).
            for v in pedidos:
                if not (v.orc and v.tipo in ("instalacao", "obra") and v.t_resposta and self._ate_hoje(v.t_envio)):
                    continue
                if self.ruido.random() >= 0.3:
                    continue
                quando = datetime.combine((v.t_envio - timedelta(days=1)).date(), time(8), tzinfo=LX)
                if quando <= v.t_resposta:
                    continue
                m, mid_t = self._msg(quando, mdm, TECNICO, f"Medição {v.orc}",
                                     f"Carlos, amanhã vais medir o {v.orc} ({v.cliente.nome}).", f"t{v.orc}", "Enviados",
                                     rng=self.ruido)
                caixa.add(m)
                m, _ = self._msg(quando + timedelta(minutes=40), TECNICO, mdm, f"Re: Medição {v.orc}",
                                 "Ok, combinado. Levo o material.\n\n> Carlos, amanhã vais medir o " + v.orc,
                                 f"t{v.orc}", "Caixa de entrada", em_resposta=mid_t, rng=self.ruido)
                caixa.add(m)
            # Ruído: newsletters, fornecedor, spam e avisos automáticos.
            d = self.inicio
            n = 0
            while d <= self.fim:
                base = datetime.combine(d, time(10), tzinfo=LX)
                if d.weekday() == 1:
                    m, _ = self._msg(base, '"Novidades Clima" <novidades@marca-exemplo.pt>', MDM_EMAIL,
                                     "Nova gama de equipamentos inverter", "Conheça a nova gama.", f"n{n}",
                                     "Caixa de entrada,Promoções", cabecalhos={"List-Unsubscribe": "<mailto:sair@marca-exemplo.pt>"})
                    caixa.add(m)
                if d.weekday() == 3:
                    m, _ = self._msg(base + timedelta(hours=2), f'"Distribuidor Exemplo" <{FORNECEDOR_EMAIL}>', MDM_EMAIL,
                                     "Fatura FT 2026/118 e orçamento de material",
                                     "Segue a fatura do material de instalação de ar condicionado e o orçamento pedido.",
                                     f"f{n}", "Caixa de entrada")
                    caixa.add(m)
                    m, _ = self._msg(base + timedelta(hours=3), f'"MDM Climatização" <{MDM_EMAIL}>', FORNECEDOR_EMAIL,
                                     "Encomenda de tubo de cobre", "Precisamos de 50 m de tubo 1/4 e 3/8.", f"f{n}", "Enviados")
                    caixa.add(m)
                if d.weekday() in (0, 4):
                    m, _ = self._msg(base + timedelta(hours=5), '"Ganhe já" <promo@spam-exemplo.biz>', MDM_EMAIL,
                                     "Orçamento grátis para o seu negócio!!!", "Clique aqui.", f"s{n}", "Spam")
                    caixa.add(m)
                if d.day == 1:
                    m, _ = self._msg(base, '"Banco Exemplo" <no-reply@banco-exemplo.pt>', MDM_EMAIL,
                                     "Extrato disponível", "O seu extrato está disponível.", f"b{n}", "Caixa de entrada")
                    caixa.add(m)
                n += 1
                d += timedelta(days=1)
        finally:
            caixa.flush()
            caixa.unlock()
            caixa.close()

    def _linha_wa(self, c: Cliente, ts: datetime, autor: str, texto: str) -> str:
        if c.formato_wa == "iphone":
            return f"‎[{ts:%d/%m/%y, %H:%M:%S}] {autor}: {texto}"
        return f"{ts:%d/%m/%y, %H:%M} - {autor}: {texto}"

    def _escrever_whatsapp(self, pasta: Path, pedidos: list[Verdade]) -> None:
        rng = self.rng
        conversas: dict[int, list[tuple[datetime, str, str]]] = {}
        nomes: dict[int, str] = {}
        for v in pedidos:
            c = v.cliente
            usa = v.canal == "whatsapp" or (v.orc and v.via_orc == "whatsapp")
            if not usa:
                continue
            if c.id not in nomes:
                nomes[c.id] = c.nome if v.canal == "whatsapp" else f"+351 {_tel_txt(c.tel)}"
            linhas = conversas.setdefault(c.id, [])
            autor = nomes[c.id]
            if v.canal == "whatsapp":
                texto = v.texto
                if rng.random() < 0.3:
                    texto += f"\nFico em {rng.choice(LOCAIS)}."
                linhas.append((v.t_contacto, autor, texto))
                # Saudação automática do WhatsApp Business: não é uma resposta.
                linhas.append((v.t_contacto + timedelta(seconds=20), MDM_WA, SAUDACAO_WA))
                for quando, texto_c in v.cobrancas:
                    if self._ate_hoje(quando):
                        linhas.append((quando, autor, texto_c))
                if v.t_resposta and self._ate_hoje(v.t_resposta):
                    linhas.append((v.t_resposta, MDM_WA, "Olá, obrigado pelo contacto! Podemos passar aí para ver. Que dia dá jeito?"))
            if v.orc and v.via_orc == "whatsapp" and self._ate_hoje(v.t_envio):
                linhas.append((v.t_envio, MDM_WA, f"Segue o orçamento {v.orc}. Qualquer dúvida diga.\n{v.orc} {c.nome}.pdf (ficheiro anexado)"))
                if v.ganho and v.aceite_por_texto and self._ate_hoje(v.t_aceite):
                    linhas.append((v.t_aceite, autor, rng.choice(ACEITES)))
                elif v.t_recusa and self._ate_hoje(v.t_recusa):
                    linhas.append((v.t_recusa, autor, RECUSAS[v.motivo]))
        por_id = {c.id: c for c in self.clientes}
        for cid, linhas in conversas.items():
            c = por_id[cid]
            linhas.sort(key=lambda x: x[0])
            primeira = linhas[0][0] - timedelta(minutes=1)
            if c.formato_wa == "iphone":
                # No iPhone, as linhas do sistema aparecem com o nome do contacto como autor.
                topo = f"[{primeira:%d/%m/%y, %H:%M:%S}] {nomes[cid]}: ‎As mensagens e as chamadas são encriptadas ponto a ponto."
                depois = linhas[-1][0] + timedelta(days=20)
                if self.ruido.random() < 0.5 and depois.date() <= self.hoje:
                    linhas.append((depois, nomes[cid], f"‎O seu código de segurança com {nomes[cid]} mudou. "
                                                       "Toque para saber mais."))
            else:
                topo = f"{primeira:%d/%m/%y, %H:%M} - As mensagens e as chamadas são encriptadas ponto a ponto. Ninguém fora desta conversa pode lê-las."
            texto = "\n".join([topo] + [self._linha_wa(c, ts, autor, t) for ts, autor, t in linhas])
            (pasta / f"Conversa do WhatsApp com {nomes[cid]}.txt").write_text(texto + "\n", encoding="utf-8")

    def _escrever_chamadas(self, caminho: Path, pedidos: list[Verdade]) -> None:
        rng = self.rng
        chamadas = []

        def acrescentar(ts, numero, tipo, duracao, nome):
            chamadas.append((ts, numero, tipo, duracao, nome))

        for v in pedidos:
            c = v.cliente
            nome = c.nome_no_telefone or "(Unknown)"
            if v.canal == "chamada":
                acrescentar(v.t_contacto, c.tel, 1 if v.atendida else 3, rng.randint(40, 400) if v.atendida else 0, nome)
                if not v.atendida and v.t_resposta and self._ate_hoje(v.t_resposta):
                    acrescentar(v.t_resposta, c.tel, 2, rng.randint(30, 300), nome)
            elif v.canal == "site" and v.t_resposta and self._ate_hoje(v.t_resposta) and (not c.email or c.id % 7 == 0):
                acrescentar(v.t_resposta, c.tel, 2, rng.randint(30, 300), nome)
        d = self.inicio
        while d <= self.fim:
            if d.weekday() == 2:
                acrescentar(datetime.combine(d, time(11, 30), tzinfo=LX), FORNECEDOR_TEL, 1, 120, "Distribuidor")
            d += timedelta(days=1)
        chamadas.sort(key=lambda x: x[0])
        linhas = ["<?xml version='1.0' encoding='UTF-8' standalone='yes' ?>", f'<calls count="{len(chamadas)}">']
        for ts, numero, tipo, duracao, nome in chamadas:
            if not self._ate_hoje(ts):
                continue
            linhas.append(f'  <call number="{numero}" duration="{duracao}" date="{int(ts.timestamp() * 1000)}" '
                          f'type="{tipo}" presentation="1" readable_date="{ts:%d/%m/%Y %H:%M:%S}" contact_name="{nome}" />')
        linhas.append("</calls>")
        caminho.write_text("\n".join(linhas) + "\n", encoding="utf-8")

    def _escrever_site(self, caminho: Path, pedidos: list[Verdade]) -> None:
        with open(caminho, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["id", "created_at", "nome", "email", "telefone", "tipo", "mensagem", "localidade"])
            for i, v in enumerate(p for p in pedidos if p.canal == "site"):
                c = v.cliente
                w.writerow([i + 1, v.t_contacto.isoformat(), c.nome, c.email or "", _tel_txt(c.tel),
                            {"instalacao": "Instalação nova", "avaria": "Reparação / avaria", "manutencao": "Manutenção",
                             "obra": "Obra de empresa"}[v.tipo], v.texto, self.rng.choice(LOCAIS)])

    def _escrever_orcamentos(self, pasta: Path, pedidos: list[Verdade]) -> None:
        from openpyxl import Workbook
        for v in pedidos:
            if not v.orc or not self._ate_hoje(v.t_envio):
                continue
            c = v.cliente
            wb = Workbook()
            ws = wb.active
            ws.title = "Orçamento"
            ws["A1"] = "MDM Climatização, Lda."
            ws["A2"] = f"Tel. 219 990 000 · {MDM_EMAIL}"
            ws["A4"], ws["B4"] = "Cliente:", c.nome
            ws["A5"], ws["B5"] = "Contacto:", (_tel_txt(c.tel) if c.id % 5 < 3 else (c.email or ""))
            ws["A6"], ws["B6"] = "Data:", datetime(v.t_envio.year, v.t_envio.month, v.t_envio.day)
            ws["B6"].number_format = "dd/mm/yyyy"
            if self.ruido.random() < 0.15:
                # Data escrita como texto, num ficheiro feito a partir de um modelo antigo.
                ws["B6"] = f"{v.t_envio:%d/%m/%Y}"
                wb.properties.created = datetime(2019, 3, 1, 9)
            ws["A8"] = f"Fornecimento e instalação de equipamento ({v.unidades} unidade(s))."
            ws["A9"], ws["A10"], ws["A11"] = "- Marca: Mitsubishi", "- Modelo: MSZ-HR", "- Cap. Nominal: Arref 3,5 kW / Aquec 3,8 kW"
            ws["A12"] = f"Valor Total: {_pt(v.valor)} € + IVA (23%) = {_pt(v.valor * 1.23)} €"
            ws["A14"] = "2ª Opção: Marca Midea"
            ws["A15"] = f"Valor Total: {_pt(v.valor * 0.9)} € + IVA (23%) = {_pt(v.valor * 0.9 * 1.23)} €"
            custos = wb.create_sheet("Custos - Mitsubishi")
            custos["A1"], custos["B1"] = "Equipamento", round(v.valor * 0.55, 2)
            custos["A2"], custos["B2"] = "Valor Total", round(v.valor * 0.8, 2)
            ano = pasta / str(v.t_envio.year)
            ano.mkdir(exist_ok=True)
            wb.save(ano / f"{v.orc} {c.nome}.xlsx")

    def _escrever_faturas(self, caminho: Path, pedidos: list[Verdade]) -> None:
        rng = self.rng
        feitas = sorted((v for v in pedidos if v.ganho and v.t_fatura and self._ate_hoje(v.t_fatura)),
                        key=lambda v: v.t_fatura)
        por_ano: dict[int, int] = {}
        descricao = {"instalacao": "Instalação de {u} equipamento(s) de ar condicionado",
                     "obra": "Climatização de obra ({u} frações)", "manutencao": "Manutenção de {u} equipamento(s)",
                     "avaria": "Reparação de avaria"}
        outros: list[list] = []
        contadores: dict[tuple[str, int], int] = {}

        def numero(serie: str, ano: int) -> str:
            contadores[(serie, ano)] = contadores.get((serie, ano), 0) + 1
            return f"{serie} {ano}/{contadores[(serie, ano)]}"

        # Pró-formas de alguns orçamentos: não são vendas e têm de ser ignoradas.
        for v in pedidos:
            if v.orc and self._ate_hoje(v.t_envio) and self.ruido.random() < 0.1:
                c = v.cliente
                outros.append([f"{v.t_envio:%d/%m/%Y}", numero("PF", v.t_envio.year), c.nome, c.nif, _pt(v.valor),
                               f"Fatura pró-forma do {v.orc}", ""])
        with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Data", "Número", "Cliente", "NIF", "Valor s/IVA", "Descrição", "Telefone"])
            for v in feitas:
                ano = v.t_fatura.year
                por_ano[ano] = por_ano.get(ano, 0) + 1
                desc = descricao[v.tipo].format(u=v.unidades)
                if v.orc and rng.random() < 0.6:
                    desc += f" (ref. {v.orc})"
                c = v.cliente
                ft = f"FT {ano}/{por_ano[ano]}"
                w.writerow([f"{v.t_fatura:%d/%m/%Y}", ft, c.nome, c.nif, _pt(v.valor), desc,
                            _tel_txt(c.tel) if rng.random() < 0.5 else ""])
                # Recibo do pagamento (não é outra venda) e, às vezes, uma nota de crédito parcial.
                recibo = v.t_fatura + timedelta(days=self.ruido.randint(0, 20))
                if recibo.date() <= self.hoje:
                    outros.append([f"{recibo:%d/%m/%Y}", numero("RG", recibo.year), c.nome, c.nif, _pt(v.valor),
                                   f"Recibo da {ft}", ""])
                credito = v.t_fatura + timedelta(days=5)
                if self.ruido.random() < 0.05 and credito.date() <= self.hoje:
                    outros.append([f"{credito:%d/%m/%Y}", numero("NC", credito.year), c.nome, c.nif,
                                   _pt(round(v.valor * 0.1, 2)), f"Desconto sobre a {ft}", ""])
            for linha in outros:
                w.writerow(linha)

    def _escrever_agenda(self, caminho: Path, pedidos: list[Verdade]) -> None:
        """A exportação da agenda do ERP: visitas para orçamentar (instalações e obras) e visitas de execução, algumas
        à espera de uma peça entre duas visitas."""
        rng, agora = self.visitas, datetime.combine(self.hoje, time(23, 59), tzinfo=LX)
        linhas: list[dict] = []

        def ate_hoje(x: datetime | None) -> datetime | None:
            return x if x is not None and x <= agora else None

        def visita(v: Verdade, tipo: str, criada: datetime, marcada: datetime | None, inicio: datetime | None,
                   feita: datetime | None, referencia: str = "", anterior: str = "", resultado: str = "concluida",
                   encomenda: date | None = None, chegada: date | None = None) -> str:
            if criada > agora:
                return ""
            if marcada is not None and inicio is not None and marcada > inicio - timedelta(hours=1):
                marcada = max(criada, inicio - timedelta(hours=1))   # marcada num fim de semana para segunda
            marcada, feita = ate_hoje(marcada), ate_hoje(feita)
            numero = str(len(linhas) + 1)
            c = v.cliente
            fmt = lambda x: x.strftime("%Y-%m-%d %H:%M:%S") if x else ""   # noqa: E731
            linhas.append({
                "visita": numero, "anterior": anterior, "cliente": c.nome, "telefone": _tel_txt(c.tel),
                "email": c.email or "", "nif": c.nif, "referencia": referencia, "tipo_trabalho": tipo,
                "prioridade": "", "tecnico": rng.choice(["Abdul", "Leonardo", "Rui"]),
                "estado": "feita" if feita else ("marcada" if marcada else "por_marcar"),
                "criada_em": fmt(criada), "marcada_em": fmt(marcada), "inicio": fmt(inicio if marcada else None),
                "fim": fmt(inicio + timedelta(hours=2) if marcada and inicio else None), "feita_em": fmt(feita),
                "cancelada_em": "", "relatorio": f"RI-{feita:%Y}-EX-{numero.zfill(4)}" if feita else "",
                "resultado": resultado if feita else "",
                "peca_encomendada_em": encomenda.isoformat() if feita and encomenda else "",
                "peca_recebida_em": chegada.isoformat() if feita and chegada and chegada <= self.hoje else ""})
            return numero

        def as_horas(dia: date, hora: int) -> datetime:
            return datetime.combine(dia, time(hora), tzinfo=LX)

        def entre(a: datetime, b: datetime, lo: float, hi: float) -> datetime:
            """Um momento em horas úteis entre a e b (a marcação faz-se no escritório)."""
            return _minuto(self.cal.somar_horas_uteis(a, self.cal.horas_uteis(a, b) * rng.uniform(lo, hi)))

        for v in sorted(pedidos, key=lambda v: v.t_contacto):
            # Visita para orçamentar: entra na agenda quando a MDM responde e acontece antes de o orçamento seguir.
            if v.tipo in ("instalacao", "obra") and v.t_resposta and v.t_envio \
                    and v.t_envio - v.t_resposta >= timedelta(days=2) and rng.random() < 0.7:
                criada = _minuto(v.t_resposta + timedelta(minutes=10))
                vao = v.t_envio - criada
                inicio = as_horas((criada + vao * rng.uniform(0.45, 0.75)).date(), rng.choice([10, 14]))
                if criada < inicio < v.t_envio - timedelta(hours=1):
                    marcada = entre(criada, inicio, 0.05, 0.4)
                    visita(v, "avaliacao", criada, marcada, inicio, inicio + timedelta(hours=1))
            # Visita de execução: entra na agenda com o sim; o trabalho acaba no dia da fatura.
            if not (v.ganho and v.t_aceite and v.t_fatura):
                continue
            criada = _minuto(v.t_aceite + timedelta(minutes=15))
            inicio = as_horas(v.t_fatura.date(), 9 if v.tipo in ("instalacao", "obra") else rng.choice([9, 14]))
            feita = inicio + timedelta(hours={"avaria": 2, "manutencao": 2, "instalacao": 6, "obra": 7}[v.tipo])
            if inicio <= criada:
                continue
            tipo = "instalacao" if v.tipo == "instalacao" else v.tipo
            if v.tipo in ("avaria", "instalacao") and inicio - criada >= timedelta(days=8) and rng.random() < 0.5:
                # Primeira visita sem a peça: fica à espera do fornecedor e volta quando ela chega.
                inicio1 = as_horas((criada + timedelta(days=rng.randint(1, 3))).date(), 10)
                feita1 = inicio1 + timedelta(hours=2)
                chegada = (feita1 + timedelta(days=rng.randint(2, max(2, min(12, (inicio - feita1).days - 2))))).date()
                if chegada >= inicio.date():
                    chegada = (inicio - timedelta(days=1)).date()
                marcada1 = entre(criada, inicio1, 0.1, 0.5)
                primeira = visita(v, tipo, criada, marcada1, inicio1, feita1, v.orc, resultado="pendente_peca",
                                  encomenda=feita1.date(), chegada=chegada)
                if primeira and feita1 <= agora:
                    marcada2 = entre(as_horas(chegada, 16), inicio, 0.0, 0.3)
                    visita(v, tipo, _minuto(feita1 + timedelta(minutes=5)), min(marcada2, inicio - timedelta(hours=1)),
                           inicio, feita, v.orc, anterior=primeira)
                continue
            marcada = entre(criada, inicio, 0.1, 0.5)
            visita(v, tipo, criada, marcada, inicio, feita, v.orc)

        with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=list(linhas[0]) if linhas else ["visita"], delimiter=";")
            w.writeheader()
            w.writerows(linhas)

    def _escrever_config(self, caminho: Path) -> None:
        caminho.write_text(f"""# Configuração do mundo de exemplo (dados inventados).
[geral]
data_referencia = "{self.hoje.isoformat()}"
exemplo = true

[empresa]
emails = ["{MDM_EMAIL}"]
telefones = ["{MDM_TEL}"]
nomes_whatsapp = ["{MDM_WA}"]

[email]
dominios_ignorar = ["distribuidor-exemplo.pt"]

[chamadas]
ignorar = ["{FORNECEDOR_TEL}"]

[faturas]
coluna_data = "Data"
coluna_numero = "Número"
coluna_cliente = "Cliente"
coluna_nif = "NIF"
coluna_valor = "Valor s/IVA"
coluna_descricao = "Descrição"
coluna_telefone = "Telefone"
coluna_email = ""
""", encoding="utf-8")

    def _escrever_verdade(self, caminho: Path) -> None:
        def f(t):
            return t.isoformat() if t else ""
        with open(caminho, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["cliente", "nome", "canal", "tipo", "empresa", "urgente", "unidades", "t_contacto", "atendida",
                        "t_resposta", "h_resposta", "orc", "t_envio", "via_orc", "valor", "ganho", "t_aceite",
                        "aceite_por_texto", "t_recusa", "motivo", "t_fatura", "cobrancas"])
            for v in self.pedidos:
                w.writerow([v.cliente.id, v.cliente.nome, v.canal, v.tipo, int(v.cliente.empresa), int(v.urgente),
                            v.unidades, f(v.t_contacto), int(v.atendida), f(v.t_resposta),
                            "" if v.h_resposta is None else f"{v.h_resposta:.4f}", v.orc, f(v.t_envio), v.via_orc,
                            "" if v.valor is None else v.valor, int(v.ganho), f(v.t_aceite), int(v.aceite_por_texto),
                            f(v.t_recusa), v.motivo, f(v.t_fatura),
                            ";".join(f(q) for q, _ in v.cobrancas if self._ate_hoje(q))])


def gerar(destino: Path | str, semente: int = 7, **opcoes) -> Mundo:
    mundo = Mundo(semente=semente, **opcoes)
    mundo.simular()
    mundo.escrever(Path(destino))
    return mundo
