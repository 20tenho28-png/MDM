"""Leitura de tabelas CSV ou Excel com cabeçalhos normalizados."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path

from ..texto import norm


def _texto_csv(caminho: Path) -> str:
    dados = caminho.read_bytes()
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return dados.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return dados.decode("latin-1", errors="replace")


def _normalizar(linha: dict) -> dict:
    return {norm(k): (v.strip() if isinstance(v, str) else v) for k, v in linha.items() if k is not None}


def ler_tabela(caminho: Path) -> list[dict]:
    """Lê .csv, .tsv, .txt, .xlsx ou .json (lista de objetos). As chaves ficam sem acentos e em minúsculas."""
    caminho = Path(caminho)
    sufixo = caminho.suffix.lower()
    if sufixo in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(caminho, data_only=True, read_only=True)
        try:
            linhas = list(wb.worksheets[0].iter_rows(values_only=True))
        finally:
            wb.close()
        if not linhas:
            return []
        cabecalho = [norm(c) for c in linhas[0]]
        return [{cabecalho[i]: v for i, v in enumerate(l) if i < len(cabecalho) and cabecalho[i]}
                for l in linhas[1:] if any(v not in (None, "") for v in l)]
    if sufixo == ".json":
        dados = json.loads(_texto_csv(caminho))
        if isinstance(dados, dict):
            dados = next((v for v in dados.values() if isinstance(v, list)), [])
        return [_normalizar(l) for l in dados if isinstance(l, dict)]
    texto = _texto_csv(caminho)
    # O separador é o que mais aparece no cabeçalho (o Excel português usa ";"; em empate, ";" ganha).
    cabecalho = texto.lstrip("﻿").splitlines()[0] if texto.strip() else ""
    contagens = {d: cabecalho.count(d) for d in (";", "\t", "|", ",")}
    delimitador = max(contagens, key=lambda d: (contagens[d], d == ";"))
    if contagens[delimitador] == 0:
        delimitador = ";"
    return [_normalizar(l) for l in csv.DictReader(io.StringIO(texto), delimiter=delimitador)]


def coluna(linha: dict, nome: str):
    if not nome:
        return ""
    valor = linha.get(norm(nome), "")
    return "" if valor is None else valor
