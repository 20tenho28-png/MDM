"""Leitura dos ficheiros de orçamento ORC (Excel), com o nome no formato 'ORC0034_2026 Nome do cliente.xlsx'."""

from __future__ import annotations

import re
from collections import Counter
from datetime import datetime, time
from pathlib import Path

from ..datas import LX, ler_data
from ..modelo import Evento
from ..texto import chave_nome, extrair_emails, extrair_telefones, norm, numero, telefone_norm

RE_FICHEIRO = re.compile(r"^\s*ORC\s*0*(\d{1,5})\s*[_\-\s]\s*(20\d{2})\s*[-_\s]*(.*?)\s*$", re.I)
_RE_TOTAL = re.compile(r"valor total\s*:?\s*(\d[\d.\s]*(?:,\d{1,2})?)\s*(?:€|eur)?\s*"
                       r"(\+\s*iva|c/\s*iva|com iva|s/\s*iva|sem iva)?")


def _valor_do_texto(texto: str) -> float | None:
    m = _RE_TOTAL.search(norm(texto))
    if not m:
        return None
    valor = numero(m.group(1))
    if valor is None:
        return None
    if (m.group(2) or "").replace(" ", "") in ("c/iva", "comiva"):
        valor = valor / 1.23
    return valor


def _numerico(x) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool) and x > 0


_RE_DATA_TEXTO = re.compile(r"\b(\d{1,2}[/.\-]\d{1,2}[/.\-]\d{4}|\d{4}-\d{2}-\d{2})\b")


def _ano_plausivel(d: datetime | None, ano: int) -> bool:
    """A data tem de ser do ano do número ORC (ou dos dias à volta da passagem de ano)."""
    if d is None:
        return False
    return d.year == ano or (d.year == ano - 1 and d.month == 12) or (d.year == ano + 1 and d.month == 1)


def ler_orcamento(caminho: Path, cfg: dict, estat: Counter) -> Evento | None:
    m = RE_FICHEIRO.match(caminho.stem)
    if not m:
        estat["nome fora do padrão ORC"] += 1
        return None
    numero_orc = f"ORC{int(m.group(1)):04d}_{m.group(2)}"
    ano = int(m.group(2))
    cliente = m.group(3).strip(" -_")
    proprios_email = {e.lower() for e in cfg["empresa"]["emails"]}
    proprios_tel = {t for t in (telefone_norm(x) for x in cfg["empresa"]["telefones"]) if t}
    folha_custos = norm(cfg["orc"]["folhas_custos"])
    try:
        from openpyxl import load_workbook
        wb = load_workbook(caminho, data_only=True, read_only=True)
    except Exception:
        estat["ilegíveis (guarda como .xlsx)"] += 1
        return None
    data, valor, textos = None, None, []
    try:
        for folha in wb.worksheets:
            if folha_custos and folha_custos in norm(folha.title):
                continue
            grelha = [list(l) for l in folha.iter_rows(max_row=300, max_col=40, values_only=True)]
            for i, linha in enumerate(grelha):
                for j, v in enumerate(linha):
                    if v is None:
                        continue
                    if isinstance(v, datetime):
                        if data is None and _ano_plausivel(v, ano):
                            data = v
                        continue
                    if not isinstance(v, str):
                        continue
                    textos.append(v)
                    if data is None and (md := _RE_DATA_TEXTO.search(v)):
                        lida = ler_data(md.group(1))
                        if lida and _ano_plausivel(lida, ano):
                            data = lida.replace(tzinfo=None)
                    if valor is None and "valor total" in norm(v):
                        valor = _valor_do_texto(v)
                        if valor is None:
                            valor = next((float(x) for x in linha[j + 1:] if _numerico(x)), None)
                        if valor is None and i + 1 < len(grelha) and j < len(grelha[i + 1]) and _numerico(grelha[i + 1][j]):
                            valor = float(grelha[i + 1][j])
        propriedades = getattr(wb, "properties", None)
    finally:
        wb.close()
    fonte_data = "celula"
    if data is None:
        candidatas = [getattr(propriedades, "created", None), getattr(propriedades, "modified", None),
                      datetime.fromtimestamp(caminho.stat().st_mtime)]
        data = next((d for d in candidatas if _ano_plausivel(d, ano)), None)
        fonte_data = "propriedades" if data else "nenhuma"
    if valor is None:
        estat["sem valor total"] += 1
    texto = "\n".join(textos)
    chaves = {f"orc:{numero_orc}"}
    k = chave_nome(cliente)
    if k:
        chaves.add(f"nome:{k}")
    chaves |= {f"email:{e}" for e in extrair_emails(texto) if e not in proprios_email}
    chaves |= {f"tel:{t}" for t in extrair_telefones(texto) if t not in proprios_tel}
    estat["lidos"] += 1
    # Sem data válida, fica provisoriamente a 1 de julho do ano do ORC; a leitura geral substitui-a pela data da
    # primeira mensagem que menciona o orçamento ou exclui o ORC (ver fontes.ler_tudo).
    ts = datetime.combine(data.date() if data else datetime(ano, 7, 1).date(), time(18), tzinfo=LX)
    return Evento(ts, "orc", None, "orcamento", f"orc:{numero_orc}",
                  nome=cliente, chaves=chaves, texto=texto, refs_orc={numero_orc}, valor=valor, numero=numero_orc,
                  extra={"ficheiro": caminho.name, "fonte_data": fonte_data})
